"""Dump every cell (formula + cached value + number format) of a workbook to stdout.

Usage: python reference/dump_workbook.py path/to/workbook.xlsx > cell_dump.txt
"""

from __future__ import annotations

import sys

import openpyxl


def main(path: str) -> None:
    wb_formulas = openpyxl.load_workbook(path, data_only=False)
    wb_values = openpyxl.load_workbook(path, data_only=True)
    for ws in wb_formulas.worksheets:
        wv = wb_values[ws.title]
        print(f"\n===== SHEET: {ws.title!r} dims={ws.dimensions}")
        if ws.merged_cells.ranges:
            print("merged:", [str(r) for r in ws.merged_cells.ranges])
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                f = c.value
                if isinstance(f, str) and f.startswith("="):
                    v = wv[c.coordinate].value
                    print(f"{c.coordinate:>6} | F: {f} | V: {v!r} | fmt={c.number_format}")
                else:
                    print(f"{c.coordinate:>6} | {f!r} | fmt={c.number_format}")


if __name__ == "__main__":
    main(sys.argv[1])
